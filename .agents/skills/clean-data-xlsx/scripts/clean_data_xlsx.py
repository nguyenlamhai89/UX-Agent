#!/usr/bin/env python3
"""Conservatively clean one XLSX workbook and publish auditable artifacts."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import uuid
from copy import copy
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl import Workbook, load_workbook


SCRIPT_VERSION = "1.0.0"
REPORT_SHEETS = [
    "Summary", "Sheet Statistics", "Missing Values", "Duplicate Rows", "Type Issues",
    "Outliers", "Merge Issues", "Cleaning Actions", "Warnings",
]
IDENTIFIER_RE = re.compile(r"(^0\d+$)|(^\d{8,}$)|([a-z].*\d|\d.*[a-z])", re.IGNORECASE)
IDENTIFIER_HEADER_RE = re.compile(r"\b(id|code|sku|zip|postal|phone|account|mã|số)\b", re.IGNORECASE)


class CleanDataError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def safe_text(value: str, header: bool = False, identifier: bool = False) -> str:
    text = value.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    if header:
        return re.sub(r"\s+", " ", text).strip()
    return text.strip() if not identifier else text


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def value_kind(value: Any) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "blank"
    if is_formula(value):
        return "formula"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    return "text"


def quantile(values: list[float], q: float) -> float:
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    position = (len(ordered) - 1) * q
    low, high = math.floor(position), math.ceil(position)
    return ordered[low] + (ordered[high] - ordered[low]) * (position - low)


def write_rows(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    for column in ws.columns:
        width = min(50, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width


def header_row(ws) -> tuple[int | None, list[str] | None]:
    for index in range(1, ws.max_row + 1):
        values = [ws.cell(index, col).value for col in range(1, ws.max_column + 1)]
        nonblank = [value for value in values if value not in (None, "")]
        if not nonblank:
            continue
        if not all(isinstance(value, str) and not is_formula(value) for value in nonblank):
            return None, None
        normalized = [safe_text(value, header=True) for value in nonblank]
        if any(not value for value in normalized) or len(set(normalized)) != len(normalized):
            return None, None
        return index, normalized
    return None, None


def row_signature(ws, row: int) -> tuple[Any, ...]:
    return tuple(ws.cell(row, col).value for col in range(1, ws.max_column + 1))


def inspect_and_clean(ws) -> dict[str, Any]:
    result: dict[str, Any] = {
        "sheet": ws.title, "before_rows": ws.max_row, "before_columns": ws.max_column,
        "after_rows": ws.max_row, "after_columns": ws.max_column, "header_row": None,
        "uncertain_header": False, "formulas": 0, "merged_ranges": [str(item) for item in ws.merged_cells.ranges],
        "normalized_headers": 0, "normalized_text_cells": 0, "removed_duplicates": 0,
        "missing": [], "type_issues": [], "outliers": [], "changes": [], "warnings": [],
    }
    formula_count = sum(1 for row in ws.iter_rows() for cell in row if is_formula(cell.value))
    result["formulas"] = formula_count
    header_index, _ = header_row(ws)
    if header_index is None:
        result["uncertain_header"] = True
        result["warnings"].append(["UNCERTAIN_HEADER", "warning", ws.title, "", "Sheet kept unchanged because its header is not safely identifiable."])
        return result
    result["header_row"] = header_index

    headers: list[str] = []
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(header_index, col)
        original = cell.value
        # A safely identified header may still contain blank cells between
        # populated columns. Preserve those cells and use a report-only label
        # so later profiling never treats None as header text.
        if original is None:
            headers.append(f"Column {col}")
            continue
        cleaned = safe_text(original, header=True)
        headers.append(cleaned)
        if cleaned != original:
            cell.value = cleaned
            result["normalized_headers"] += 1
            result["changes"].append([ws.title, cell.coordinate, "HEADER_WHITESPACE_NORMALIZATION", original, cleaned])

    kinds: list[set[str]] = [set() for _ in range(ws.max_column)]
    numbers: list[list[float]] = [[] for _ in range(ws.max_column)]
    missing: list[int] = [0 for _ in range(ws.max_column)]
    for row in range(header_index + 1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            kind = value_kind(cell.value)
            if kind == "blank":
                missing[col - 1] += 1
                continue
            kinds[col - 1].add(kind)
            if kind == "number":
                numbers[col - 1].append(float(cell.value))
            if kind == "text":
                identifier = bool(IDENTIFIER_RE.search(cell.value) or IDENTIFIER_HEADER_RE.search(headers[col - 1]))
                cleaned = safe_text(cell.value, identifier=identifier)
                if cleaned != cell.value:
                    original = cell.value
                    cell.value = cleaned
                    result["normalized_text_cells"] += 1
                    result["changes"].append([ws.title, cell.coordinate, "DATA_WHITESPACE_NORMALIZATION", original, cleaned])

    for col, count in enumerate(missing, start=1):
        result["missing"].append([ws.title, headers[col - 1], count, ws.max_row - header_index])
    for col, found in enumerate(kinds, start=1):
        data_kinds = sorted(found - {"formula"})
        if len(data_kinds) > 1:
            result["type_issues"].append([ws.title, headers[col - 1], ", ".join(data_kinds), "not_converted"])
            result["warnings"].append(["MIXED_TYPES_REVIEW", "warning", ws.title, headers[col - 1], "Mixed data types were retained."])
    for col, values in enumerate(numbers, start=1):
        if len(values) < 4:
            continue
        q1, q3 = quantile(values, 0.25), quantile(values, 0.75)
        iqr = q3 - q1
        low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        flagged = sum(1 for value in values if value < low or value > high)
        result["outliers"].append([ws.title, headers[col - 1], len(values), q1, q3, low, high, flagged, "flagged_not_modified"])

    if result["merged_ranges"]:
        result["warnings"].append(["MERGED_CELLS_REVIEW", "warning", ws.title, ", ".join(result["merged_ranges"]), "Merged cells were preserved; duplicate removal skipped."])
    if formula_count:
        result["warnings"].append(["FORMULAS_PRESERVED", "info", ws.title, "", "Formula cells were not modified; duplicate removal skipped."])
    if not result["merged_ranges"] and not formula_count:
        seen: dict[tuple[Any, ...], int] = {}
        duplicates: list[int] = []
        for row in range(header_index + 1, ws.max_row + 1):
            signature = row_signature(ws, row)
            if all(value in (None, "") for value in signature):
                continue
            if signature in seen:
                duplicates.append(row)
                result["changes"].append([ws.title, f"row:{row}", "EXACT_DUPLICATE_ROW_REMOVAL", f"duplicate_of_row:{seen[signature]}", "removed"])
            else:
                seen[signature] = row
        for row in reversed(duplicates):
            ws.delete_rows(row, 1)
        result["removed_duplicates"] = len(duplicates)
    result["after_rows"] = ws.max_row
    return result


def build_report(path: Path, sheet_results: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    report = Workbook()
    report.remove(report.active)
    sheets = {name: report.create_sheet(name) for name in REPORT_SHEETS}
    write_rows(sheets["Summary"], ["Metric", "Value"], [[key, value] for key, value in summary.items()])
    write_rows(sheets["Sheet Statistics"], ["Sheet", "Rows Before", "Rows After", "Columns", "Header Row", "Headers Normalized", "Text Cells Normalized", "Duplicates Removed", "Formulas", "Merged Ranges", "Review Status"], [
        [item["sheet"], item["before_rows"], item["after_rows"], item["after_columns"], item["header_row"], item["normalized_headers"], item["normalized_text_cells"], item["removed_duplicates"], item["formulas"], len(item["merged_ranges"]), "review" if item["warnings"] else "ok"] for item in sheet_results])
    write_rows(sheets["Missing Values"], ["Sheet", "Column", "Missing Count", "Data Rows"], [row for item in sheet_results for row in item["missing"]])
    write_rows(sheets["Duplicate Rows"], ["Sheet", "Exact Duplicates Removed", "Rule"], [[item["sheet"], item["removed_duplicates"], "exact_only"] for item in sheet_results])
    write_rows(sheets["Type Issues"], ["Sheet", "Column", "Kinds", "Action"], [row for item in sheet_results for row in item["type_issues"]])
    write_rows(sheets["Outliers"], ["Sheet", "Column", "Numeric Values", "Q1", "Q3", "Lower Bound", "Upper Bound", "Flagged", "Action"], [row for item in sheet_results for row in item["outliers"]])
    merge_rows = [[item["sheet"], merged, "preserved", "merged_cell_structure"] for item in sheet_results for merged in item["merged_ranges"]]
    merge_rows.append(["workbook", "not_applicable", "0", "relational joins: no explicit key contract"])
    write_rows(sheets["Merge Issues"], ["Sheet", "Range/Status", "Action", "Notes"], merge_rows)
    write_rows(sheets["Cleaning Actions"], ["Sheet", "Location", "Rule", "Before", "After"], [row for item in sheet_results for row in item["changes"]])
    write_rows(sheets["Warnings"], ["Code", "Severity", "Sheet", "Location", "Message"], [row for item in sheet_results for row in item["warnings"]])
    report.save(path)


def existing_output_is_valid(root: Path, source_name: str, source_hash: str) -> bool:
    analysis = root / "Analysis"
    log_path = analysis / "cleaning_log.json"
    raw_path = analysis / source_name
    required_derivatives = [
        root / f"{Path(source_name).stem}_cleaned.xlsx",
        root / "Scripts" / "clean.py",
        analysis / "data_quality_report.xlsx",
    ]
    if not log_path.is_file() or not raw_path.is_file() or not all(path.is_file() for path in required_derivatives):
        return False
    try:
        log = json.loads(log_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    return log.get("input", {}).get("source_sha256") == source_hash and sha256(raw_path) == source_hash


def stage_paths(stage_root: Path, stem: str, source_name: str) -> dict[str, Path]:
    return {
        "cleaned": stage_root / f"{stem}_cleaned.xlsx",
        "script": stage_root / "Scripts" / "clean.py",
        "raw": stage_root / "Analysis" / source_name,
        "report": stage_root / "Analysis" / "data_quality_report.xlsx",
        "log": stage_root / "Analysis" / "cleaning_log.json",
    }


def publish_new(stage_root: Path, root: Path) -> None:
    os.replace(stage_root, root)


def publish_refresh(stage: dict[str, Path], root: Path, stem: str, source_name: str) -> None:
    final = stage_paths(root, stem, source_name)
    targets = ["cleaned", "script", "report", "log"]
    backups: list[tuple[Path, Path]] = []
    token = uuid.uuid4().hex
    try:
        for name in targets:
            target = final[name]
            backup = target.with_name(target.name + f".backup-{token}")
            if target.exists():
                os.replace(target, backup)
                backups.append((backup, target))
            os.replace(stage[name], target)
        for backup, _ in backups:
            backup.unlink(missing_ok=True)
    except Exception:
        for _, target in reversed(backups):
            backup = next((item[0] for item in backups if item[1] == target), None)
            if backup and backup.exists():
                os.replace(backup, target)
        raise


def clean_workbook(input_path: str | Path, replay_root: Path | None = None, expected_hash: str | None = None) -> dict[str, Any]:
    source = Path(input_path).expanduser().resolve()
    if not source.exists():
        raise CleanDataError("INPUT_NOT_FOUND", f"Input file not found: {source}")
    if not source.is_file():
        raise CleanDataError("INPUT_NOT_FILE", f"Input path is not a file: {source}")
    if source.suffix.lower() != ".xlsx":
        raise CleanDataError("UNSUPPORTED_EXTENSION", "Only .xlsx files are supported.")
    source_hash = sha256(source)
    if expected_hash and source_hash != expected_hash:
        raise CleanDataError("RAW_ARCHIVE_HASH_MISMATCH", "Raw archive hash does not match the recorded source hash.")
    root = replay_root.resolve() if replay_root else source.parent / source.stem
    parent = root.parent
    if not os.access(parent, os.W_OK):
        raise CleanDataError("OUTPUT_PARENT_UNWRITABLE", f"Cannot write beside input file: {parent}")
    existing = root.exists()
    if existing and not existing_output_is_valid(root, source.name, source_hash):
        code = "OUTPUT_COLLISION_RAW_MISMATCH" if (root / "Analysis" / source.name).exists() else "OUTPUT_COLLISION_INVALID_STATE"
        raise CleanDataError(code, f"Existing output cannot be safely refreshed: {root}")
    try:
        workbook = load_workbook(source, data_only=False)
    except Exception as error:
        raise CleanDataError("INVALID_XLSX", f"Unable to open XLSX safely: {error}") from error
    stage_root = Path(tempfile.mkdtemp(prefix=f".{root.name}.clean-data-", dir=parent)) / root.name
    try:
        paths = stage_paths(stage_root, source.stem, source.name)
        for path in paths.values():
            path.parent.mkdir(parents=True, exist_ok=True)
        if not existing:
            shutil.copyfile(source, paths["raw"])
            if sha256(paths["raw"]) != source_hash:
                raise CleanDataError("RAW_ARCHIVE_HASH_MISMATCH", "Raw archive copy verification failed.")
        sheet_results = [inspect_and_clean(sheet) for sheet in workbook.worksheets]
        workbook.save(paths["cleaned"])
        try:
            validation_book = load_workbook(paths["cleaned"], data_only=False)
            if [sheet.title for sheet in validation_book.worksheets] != [sheet.title for sheet in workbook.worksheets]:
                raise ValueError("worksheet names/order changed")
        except Exception as error:
            raise CleanDataError("POST_WRITE_VALIDATION_FAILED", str(error)) from error
        warnings = [row for item in sheet_results for row in item["warnings"]]
        summary = {
            "status": "success", "source_file": str(source), "output_dir": str(root),
            "source_sha256": source_hash, "worksheets": len(sheet_results),
            "duplicates_removed": sum(item["removed_duplicates"] for item in sheet_results),
            "warnings": len(warnings), "join_analysis": "not_applicable",
        }
        build_report(paths["report"], sheet_results, summary)
        shutil.copyfile(Path(__file__).resolve(), paths["script"])
        script_hash = sha256(paths["script"])
        final_paths = stage_paths(root, source.stem, source.name)
        log = {
            "schema_version": "1.0", "status": "success",
            "run": {"tool_name": "clean-data-xlsx", "tool_version": SCRIPT_VERSION, "replay_mode": replay_root is not None, "library_versions": {"openpyxl": openpyxl.__version__, "pandas": pd.__version__}},
            "input": {"source_path": str(source), "source_filename": source.name, "source_sha256": source_hash, "raw_archive_path": str(final_paths["raw"]), "raw_archive_sha256": source_hash, "raw_archive_verified": True},
            "output": {"output_dir": str(root), "cleaned_xlsx_path": str(final_paths["cleaned"]), "report_xlsx_path": str(final_paths["report"])},
            "packaged_replay_script": {"path": str(final_paths["script"]), "script_version": SCRIPT_VERSION, "sha256": script_hash, "raw_input_resolution": f"../Analysis/{source.name}", "api_key_required": False},
            "rules": ["HEADER_WHITESPACE_NORMALIZATION", "DATA_WHITESPACE_NORMALIZATION", "IDENTIFIER_PRESERVATION", "FORMULA_PRESERVATION", "EXACT_DUPLICATE_ROW_REMOVAL", "MISSING_VALUE_NO_IMPUTATION", "OUTLIER_FLAG_ONLY", "TYPE_CONVERSION_DISABLED", "NO_CROSS_SHEET_JOIN"],
            "workbook_inspection": {"sheets": sheet_results, "relational_merge_analysis": {"status": "not_applicable", "joins_performed": 0}},
            "changes": [row for item in sheet_results for row in item["changes"]],
            "warnings": warnings,
            "validation": {"cleaned_workbook_reopened": True, "raw_sha256_verified": True, "report_generated": True, "replay_script_generated": True},
            "errors": [],
        }
        paths["log"].write_text(json.dumps(log, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        if existing:
            publish_refresh(paths, root, source.stem, source.name)
        else:
            publish_new(stage_root, root)
        result = {"status": "success", "input_xlsx_path": str(source), "output_dir": str(root), "cleaned_xlsx_path": str(final_paths["cleaned"]), "raw_xlsx_path": str(final_paths["raw"]), "report_xlsx_path": str(final_paths["report"]), "cleaning_log_path": str(final_paths["log"]), "replay_script_path": str(final_paths["script"]), "raw_sha256": source_hash, "cleaned_sha256": sha256(final_paths["cleaned"]), "script_sha256": script_hash, "warnings": [{"code": row[0], "sheet": row[2], "message": row[4]} for row in warnings]}
        return result
    finally:
        stage_parent = stage_root.parent
        if stage_parent.exists():
            shutil.rmtree(stage_parent, ignore_errors=True)


def replay_from_self(script_path: Path) -> dict[str, Any]:
    root = script_path.resolve().parent.parent
    log_path = root / "Analysis" / "cleaning_log.json"
    if not log_path.is_file():
        raise CleanDataError("REPLAY_RAW_INPUT_MISSING", "Replay log not found beside clean.py.")
    log = json.loads(log_path.read_text(encoding="utf-8"))
    source_name = log.get("input", {}).get("source_filename")
    expected_hash = log.get("input", {}).get("raw_archive_sha256")
    raw = root / "Analysis" / str(source_name)
    if not raw.is_file():
        raise CleanDataError("REPLAY_RAW_INPUT_MISSING", f"Raw archive not found: {raw}")
    return clean_workbook(raw, replay_root=root, expected_hash=expected_hash)


def main(argv: list[str] | None = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    try:
        if not args and Path(__file__).name == "clean.py":
            result = replay_from_self(Path(__file__))
        else:
            parser = argparse.ArgumentParser(description="Clean one XLSX workbook safely.")
            parser.add_argument("input_xlsx_path", nargs="?", help="Path to one .xlsx file")
            parsed = parser.parse_args(args)
            if not parsed.input_xlsx_path:
                raise CleanDataError("INPUT_PATH_REQUIRED", "Provide exactly one .xlsx path.")
            result = clean_workbook(parsed.input_xlsx_path)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except CleanDataError as error:
        print(json.dumps({"status": "error", "code": error.code, "message": error.message}, ensure_ascii=False), file=sys.stderr)
        return 1
    except Exception as error:
        print(json.dumps({"status": "error", "code": "INTERNAL_ERROR", "message": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
