from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    import pytest
except ImportError:
    class _MockPytestRaises:
        def __init__(self, expected_exception):
            self.expected_exception = expected_exception
            self.value = None

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc_val, exc_tb):
            if exc_type is None:
                raise AssertionError(f"Expected {self.expected_exception.__name__} but no exception was raised.")
            if issubclass(exc_type, self.expected_exception):
                self.value = exc_val
                return True
            return False

    class _MockPytestMark:
        @staticmethod
        def parametrize(*args, **kwargs):
            def decorator(func):
                return func
            return decorator

    class _MockPytest:
        raises = _MockPytestRaises
        mark = _MockPytestMark

        @staticmethod
        def fixture(func):
            return func

    pytest = _MockPytest()
from openpyxl import Workbook, load_workbook



SCRIPT_PATH = Path(__file__).parents[1] / "scripts" / "clean_data_xlsx.py"
SPEC = importlib.util.spec_from_file_location("clean_data_xlsx", SCRIPT_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)

WORKSPACE = Path(__file__).parents[4]
TEST_BASE = WORKSPACE / ".agents" / ".test-tmp" / "clean-data-xlsx"


@pytest.fixture
def case_dir():
    TEST_BASE.mkdir(parents=True, exist_ok=True)
    created = Path(tempfile.mkdtemp(prefix="case-", dir=TEST_BASE))
    try:
        yield created
    finally:
        shutil.rmtree(created, ignore_errors=True)


def create_workbook(path: Path, *, formula: bool = False, numeric_header: bool = False, merged: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Customers"
    if numeric_header:
        sheet.append([101, 202])
    else:
        sheet.append([" Customer\u00a0 Name ", "Customer ID", "Revenue"])
        sheet.append([" Alice ", "0012", 1])
        sheet.append([" Alice ", "0012", 1])
        sheet.append(["", "0042", 2])
        sheet.append(["Bob", "0099", 3])
        sheet.append(["Cara", "0100", 100])
        if formula:
            sheet.cell(7, 1, "Formula row")
            sheet.cell(7, 2, "0200")
            sheet.cell(7, 3, "=SUM(C2:C6)")
            sheet.cell(8, 1, "Formula row")
            sheet.cell(8, 2, "0200")
            sheet.cell(8, 3, "=SUM(C2:C6)")
        if merged:
            sheet.merge_cells("A9:B9")
            sheet["A9"] = "Merged note"
    second = workbook.create_sheet("Other")
    second.append([" Code ", "Value"])
    second.append(["007", " text "])
    workbook.save(path)


def output_root(source: Path) -> Path:
    return source.parent / source.stem


def test_creates_required_tree_preserves_raw_and_applies_safe_cleaning(case_dir: Path):
    source = case_dir / "Customer Data.xlsx"
    create_workbook(source)
    source_bytes = source.read_bytes()

    result = MODULE.clean_workbook(source)
    root = output_root(source)

    assert result["status"] == "success"
    assert source.read_bytes() == source_bytes
    assert (root / "Customer Data_cleaned.xlsx").is_file()
    assert (root / "Scripts" / "clean.py").is_file()
    assert (root / "Analysis" / "Customer Data.xlsx").read_bytes() == source_bytes
    assert (root / "Analysis" / "data_quality_report.xlsx").is_file()
    log = json.loads((root / "Analysis" / "cleaning_log.json").read_text(encoding="utf-8"))
    assert log["input"]["raw_archive_verified"] is True
    assert log["packaged_replay_script"]["raw_input_resolution"] == "../Analysis/Customer Data.xlsx"
    cleaned = load_workbook(root / "Customer Data_cleaned.xlsx", data_only=False)
    sheet = cleaned["Customers"]
    assert [sheet.cell(1, column).value for column in range(1, 4)] == ["Customer Name", "Customer ID", "Revenue"]
    assert sheet["B2"].value == "0012"
    assert sheet.max_row == 5
    assert sheet["C5"].value == 100
    report = load_workbook(root / "Analysis" / "data_quality_report.xlsx", data_only=True)
    assert report.sheetnames == MODULE.REPORT_SHEETS
    assert report["Outliers"].max_row > 1


def test_formula_rows_are_preserved_and_deduplication_is_skipped(case_dir: Path):
    source = case_dir / "Formula.xlsx"
    create_workbook(source, formula=True)

    MODULE.clean_workbook(source)
    sheet = load_workbook(output_root(source) / "Formula_cleaned.xlsx", data_only=False)["Customers"]
    assert sheet["C7"].value == "=SUM(C2:C6)"
    assert sheet["C8"].value == "=SUM(C2:C6)"
    log = json.loads((output_root(source) / "Analysis" / "cleaning_log.json").read_text(encoding="utf-8"))
    assert any(item[0] == "FORMULAS_PRESERVED" for item in log["warnings"])


def test_merged_cells_are_reported_and_not_removed(case_dir: Path):
    source = case_dir / "Merged.xlsx"
    create_workbook(source, merged=True)

    MODULE.clean_workbook(source)
    cleaned = load_workbook(output_root(source) / "Merged_cleaned.xlsx", data_only=False)
    assert "A9:B9" in {str(item) for item in cleaned["Customers"].merged_cells.ranges}
    report = load_workbook(output_root(source) / "Analysis" / "data_quality_report.xlsx", data_only=True)
    assert report["Merge Issues"].max_row >= 3


def test_uncertain_header_stays_unchanged(case_dir: Path):
    source = case_dir / "Uncertain.xlsx"
    create_workbook(source, numeric_header=True)

    MODULE.clean_workbook(source)
    cleaned = load_workbook(output_root(source) / "Uncertain_cleaned.xlsx", data_only=False)
    assert cleaned["Customers"]["A1"].value == 101
    log = json.loads((output_root(source) / "Analysis" / "cleaning_log.json").read_text(encoding="utf-8"))
    assert any(item[0] == "UNCERTAIN_HEADER" for item in log["warnings"])


def test_replay_script_uses_raw_archive_without_creating_nested_folder(case_dir: Path):
    source = case_dir / "Replay.xlsx"
    create_workbook(source)
    MODULE.clean_workbook(source)
    root = output_root(source)

    completed = subprocess.run([sys.executable, str(root / "Scripts" / "clean.py")], text=True, capture_output=True, check=False)

    assert completed.returncode == 0, completed.stderr
    assert (root / "Replay_cleaned.xlsx").is_file()
    assert not (root / "Analysis" / "Replay").exists()
    replay_log = json.loads((root / "Analysis" / "cleaning_log.json").read_text(encoding="utf-8"))
    assert replay_log["run"]["replay_mode"] is True


def test_mismatched_existing_raw_refuses_overwrite(case_dir: Path):
    source = case_dir / "Collision.xlsx"
    create_workbook(source)
    MODULE.clean_workbook(source)
    original_cleaned = (output_root(source) / "Collision_cleaned.xlsx").read_bytes()
    create_workbook(source, formula=True)

    with pytest.raises(MODULE.CleanDataError) as error:
        MODULE.clean_workbook(source)

    assert error.value.code == "OUTPUT_COLLISION_RAW_MISMATCH"
    assert (output_root(source) / "Collision_cleaned.xlsx").read_bytes() == original_cleaned


@pytest.mark.parametrize("name", ["not_excel.csv", "not_excel.txt"])
def test_rejects_non_xlsx_input(case_dir: Path, name: str):
    source = case_dir / name
    source.write_text("data", encoding="utf-8")

    with pytest.raises(MODULE.CleanDataError) as error:
        MODULE.clean_workbook(source)

    assert error.value.code == "UNSUPPORTED_EXTENSION"


def test_cli_requires_exactly_one_input(case_dir: Path):
    completed = subprocess.run([sys.executable, str(SCRIPT_PATH)], text=True, capture_output=True, check=False)
    assert completed.returncode == 1
    assert "INPUT_PATH_REQUIRED" in completed.stderr


if __name__ == "__main__":
    try:
        import pytest
        sys.exit(pytest.main([__file__]))
    except ImportError:
        import inspect
        print("pytest not installed; running tests via standard python test runner...")
        passed = 0
        failed = 0
        test_funcs = [
            (name, func) for name, func in list(globals().items())
            if name.startswith("test_") and callable(func)
        ]
        for name, func in test_funcs:
            if name == "test_rejects_non_xlsx_input":
                sub_names = ["not_excel.csv", "not_excel.txt"]
            else:
                sub_names = [None]
            for sub in sub_names:
                TEST_BASE.mkdir(parents=True, exist_ok=True)
                c_dir = Path(tempfile.mkdtemp(prefix="case-", dir=TEST_BASE))
                try:
                    sig = inspect.signature(func)
                    params = list(sig.parameters.keys())
                    if len(params) == 2 and sub is not None:
                        func(c_dir, sub)
                    elif len(params) == 1:
                        func(c_dir)
                    else:
                        func()
                    print(f"  PASSED: {name}{'[' + sub + ']' if sub else ''}")
                    passed += 1
                except Exception as error:
                    print(f"  FAILED: {name}{'[' + sub + ']' if sub else ''}: {error}")
                    failed += 1
                finally:
                    shutil.rmtree(c_dir, ignore_errors=True)
        print(f"\nTest Summary: {passed} passed, {failed} failed.")
        sys.exit(0 if failed == 0 else 1)

